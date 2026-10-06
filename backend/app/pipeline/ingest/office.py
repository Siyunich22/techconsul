"""Офисные документы (docx, doc, rtf, odt, pptx, ppt): конвертация в PDF через LibreOffice.

PDF-версия даёт реальные номера страниц (для ссылок-источников «документ + страница») и единый
предпросмотр с подсветкой. Если LibreOffice не справился — текст docx/pptx извлекается напрямую.
"""

import logging
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from app.core.config import get_settings
from app.pipeline.ingest.base import IngestError, Page, Parsed, Table, clean_text
from app.pipeline.ingest.pdf import parse_pdf

log = logging.getLogger(__name__)


def soffice_convert(src: Path, out_dir: Path, target: str) -> Path:
    """soffice --headless --convert-to <target>.

    Отдельный профиль LibreOffice на вызов — параллельные воркеры не мешают друг другу."""
    profile = Path(tempfile.gettempdir()) / f"lo_profile_{uuid.uuid4().hex}"
    cmd = [
        "soffice",
        f"-env:UserInstallation=file://{profile.as_posix()}",
        "--headless",
        "--norestore",
        "--convert-to",
        target,
        "--outdir",
        str(out_dir),
        str(src),
    ]
    try:
        subprocess.run(cmd, capture_output=True, timeout=get_settings().soffice_timeout_s, check=True)
    except FileNotFoundError as exc:
        raise IngestError("LibreOffice не установлен") from exc
    except subprocess.TimeoutExpired as exc:
        raise IngestError("Конвертация LibreOffice превысила лимит времени") from exc
    except subprocess.CalledProcessError as exc:
        raise IngestError(f"LibreOffice: {exc.stderr.decode(errors='replace')[:300]}") from exc
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    ext = target.split(":", 1)[0]
    result = out_dir / f"{src.stem}.{ext}"
    if not result.exists():
        raise IngestError("LibreOffice не создал выходной файл")
    return result


def parse_office(path: Path, work_dir: Path) -> Parsed:
    try:
        pdf = soffice_convert(path, work_dir, "pdf")
    except IngestError as exc:
        log.warning("Конвертация %s в PDF не удалась (%s) — прямое извлечение текста", path.name, exc)
        return _direct(path)
    parsed = parse_pdf(pdf, kind="office")
    parsed.preview_pdf = pdf
    if path.suffix.lower() == ".docx":
        _attach_docx_tables(path, parsed)
    return parsed


def _direct(path: Path) -> Parsed:
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        text = clean_text("\n".join(p.text for p in d.paragraphs))
        tables = [Table(rows=[[clean_text(c.text) for c in row.cells] for row in t.rows]) for t in d.tables]
        return Parsed(
            kind="office", pages=[Page(no=1, text=text, tables=tables)], meta={"pagination": "none"}
        )
    if ext == ".pptx":
        from pptx import Presentation

        pages = []
        for i, slide in enumerate(Presentation(str(path)).slides, start=1):
            parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
            pages.append(Page(no=i, text=clean_text("\n".join(parts))))
        return Parsed(kind="office", pages=pages)
    raise IngestError(f"Не удалось обработать файл {ext}: LibreOffice недоступен")


def _attach_docx_tables(path: Path, parsed: Parsed) -> None:
    """Таблицы docx точнее из самого docx, чем из PDF: привязываем к странице по тексту первой ячейки."""
    import docx

    try:
        d = docx.Document(str(path))
    except Exception:
        return
    for t in d.tables:
        rows = [[clean_text(c.text) for c in row.cells] for row in t.rows]
        rows = [r for r in rows if any(r)]
        if len(rows) < 2:
            continue
        probe = next((c for c in rows[0] if len(c) >= 4), None)
        page = next((p for p in parsed.pages if probe and probe in p.text), None)
        if page is None:
            continue
        page.tables = [tb for tb in page.tables if tb.rows[0] != rows[0]] + [Table(rows=rows)]
