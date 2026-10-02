"""Таблицы: xlsx/xlsm (openpyxl, формулы → значения), xls (xlrd), csv.

Каждый лист — «страница» с адресами строк: чанки получают диапазон ячеек как источник (TZ §5.1).
Если в книге формулы без сохранённых значений (файл не пересчитан в Excel) — пересчёт LibreOffice.
"""

import csv
import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.pipeline.ingest.base import IngestError, Page, Parsed, Table

log = logging.getLogger(__name__)

MAX_ROWS_PER_SHEET = 20000
MAX_COLS = 60


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.6g}" if abs(v) < 1e15 else f"{v:.0f}"
    return str(v).strip()


def _sheet_page(no: int, name: str, rows: list[tuple[int, list[str]]]) -> Page:
    rows = [(r, cells) for r, cells in rows if any(cells)]
    width = max((len(c) for _, c in rows), default=0)
    width = min(width, MAX_COLS)
    lines, row_texts = [], []
    for r, cells in rows:
        cells = cells[:width]
        line = f"[{r}] " + " | ".join(cells)
        lines.append(line)
        row_texts.append((r, line))
    max_col = get_column_letter(width) if width else "A"
    table = Table(
        rows=[cells[:width] for _, cells in rows][:500],
        title=name,
        cell_range=f"A{rows[0][0]}:{max_col}{rows[-1][0]}" if rows else None,
    )
    return Page(
        no=no,
        text=f"Лист «{name}»\n" + "\n".join(lines),
        tables=[table] if rows else [],
        sheet=name,
        rows=row_texts,
        max_col=max_col,
    )


def _formulas_without_values(path: Path) -> bool:
    wb_f = load_workbook(path, read_only=True, data_only=False)
    wb_v = load_workbook(path, read_only=True, data_only=True)
    formulas = missing = 0
    try:
        for ws_f in wb_f.worksheets:
            ws_v = wb_v[ws_f.title]
            for row_f, row_v in zip(ws_f.iter_rows(max_row=2000), ws_v.iter_rows(max_row=2000), strict=False):
                for cf, cv in zip(row_f, row_v, strict=False):
                    if isinstance(cf.value, str) and cf.value.startswith("="):
                        formulas += 1
                        missing += cv.value is None
    finally:
        wb_f.close()
        wb_v.close()
    return formulas > 0 and missing / formulas > 0.5


def parse_xlsx(path: Path, work_dir: Path) -> Parsed:
    meta: dict = {}
    try:
        if _formulas_without_values(path):
            from app.pipeline.ingest.office import soffice_convert

            # конвертация копии в xlsx заставляет LibreOffice пересчитать формулы и сохранить значения
            src = work_dir / f"recalc_{path.stem}{path.suffix}"
            src.write_bytes(path.read_bytes())
            out_dir = work_dir / "recalc"
            out_dir.mkdir(exist_ok=True)
            path = soffice_convert(src, out_dir, "xlsx")
            meta["recalculated"] = True
        wb = load_workbook(path, read_only=True, data_only=True)
    except IngestError:
        raise
    except Exception as exc:
        raise IngestError(f"Не удалось открыть книгу Excel: {exc}") from exc
    pages = []
    try:
        for i, ws in enumerate(wb.worksheets, start=1):
            rows = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True, max_row=MAX_ROWS_PER_SHEET), start=1):
                rows.append((r_idx, [_fmt(v) for v in row[:MAX_COLS]]))
            pages.append(_sheet_page(i, ws.title, rows))
    finally:
        wb.close()
    meta["sheets"] = [p.sheet for p in pages]
    return Parsed(kind="sheet", pages=pages, meta=meta)


def parse_xls(path: Path) -> Parsed:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise IngestError(f"Не удалось открыть книгу Excel (xls): {exc}") from exc
    pages = []
    for i, sh in enumerate(book.sheets(), start=1):
        rows = [
            (r + 1, [_fmt(v) for v in sh.row_values(r)[:MAX_COLS]])
            for r in range(min(sh.nrows, MAX_ROWS_PER_SHEET))
        ]
        pages.append(_sheet_page(i, sh.name, rows))
    return Parsed(kind="sheet", pages=pages, meta={"sheets": [p.sheet for p in pages]})


def parse_csv(path: Path) -> Parsed:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [
        (i, [c.strip() for c in row]) for i, row in enumerate(csv.reader(text.splitlines(), dialect), start=1)
    ]
    return Parsed(kind="sheet", pages=[_sheet_page(1, path.stem, rows[:MAX_ROWS_PER_SHEET])])
